from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.services.admin.admin_analytics_service import (
    AdminAnalyticsService
)


class AdminReportExportService:

    def __init__(self):

        self.analytics = (
            AdminAnalyticsService()
        )

    # ====================================
    # AUTO RESIZE SHEET COLUMNS
    # ====================================

    def _auto_resize_columns(self, sheet):

        for column in sheet.columns:

            max_length = 0

            column_letter = get_column_letter(
                column[0].column
            )

            for cell in column:

                try:

                    if cell.value:

                        max_length = max(
                            max_length,
                            len(str(cell.value))
                        )

                except Exception:
                    pass

            sheet.column_dimensions[
                column_letter
            ].width = max_length + 5

    # ====================================
    # GENERATE EXCEL REPORT
    # ====================================

    def generate_excel_report(self):

        workbook = Workbook()

        # ====================================
        # PLATFORM SUMMARY
        # ====================================

        summary_sheet = workbook.active

        summary_sheet.title = (
            "Platform Summary"
        )

        stats = (
            self.analytics
            .get_platform_statistics()
        )

        summary_sheet.append(
            ["Metric", "Value"]
        )

        for key, value in stats.items():

            summary_sheet.append(
                [key, value]
            )

        self._auto_resize_columns(
            summary_sheet
        )

        # ====================================
        # USER GROWTH
        # ====================================

        growth_sheet = (
            workbook.create_sheet(
                "User Growth"
            )
        )

        growth_sheet.append(
            ["Date", "Users"]
        )

        for row in (
            self.analytics
            .get_user_growth()
        ):

            growth_sheet.append([
                row["date"],
                row["count"]
            ])

        self._auto_resize_columns(
            growth_sheet
        )

        # ====================================
        # ACCURACY TRENDS
        # ====================================

        accuracy_sheet = (
            workbook.create_sheet(
                "Accuracy Trends"
            )
        )

        accuracy_sheet.append(
            ["Date", "Accuracy"]
        )

        for row in (
            self.analytics
            .get_accuracy_trends()
        ):

            accuracy_sheet.append([
                row["date"],
                row["accuracy"]
            ])

        self._auto_resize_columns(
            accuracy_sheet
        )

        # ====================================
        # ASSESSMENT ACTIVITY
        # ====================================

        activity_sheet = (
            workbook.create_sheet(
                "Assessment Activity"
            )
        )

        activity_sheet.append(
            ["Date", "Assessments"]
        )

        for row in (
            self.analytics
            .get_assessment_activity()
        ):

            activity_sheet.append([
                row["date"],
                row["assessments"]
            ])

        self._auto_resize_columns(
            activity_sheet
        )

        # ====================================
        # SAVE FILE
        # ====================================

        output = BytesIO()

        workbook.save(output)

        output.seek(0)

        return output